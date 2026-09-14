import io
import datetime
import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
import streamlit as st
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

# --- SAYFA YAPILANDIRMASI ---
st.set_page_config(page_title="Tabur Veri Yönetim ve Raporlama Sistemi", layout="wide")

# --- MOCK VERİ TABANI & OTURUM YÖNETİMİ ---
if "users_db" not in st.session_state:
    st.session_state["users_db"] = {
        "admin": {"password": "123", "role": "admin", "name": "Sistem Yöneticisi"},
        "reporter": {"password": "123", "role": "reporter", "name": "Raporlama Personeli"},
    }

if "logged_in_user" not in st.session_state:
    st.session_state["logged_in_user"] = None

# Örnek Veri Seti (Tabur / Bölük / Tim / Personel Hiyerarşisi)
if "main_data" not in st.session_state:
    st.session_state["main_data"] = pd.DataFrame([
        {"Personel": "Ahmet Yılmaz", "Tabur": "1. Tabur", "Bölük": "1. Bölük", "Tim": "1. Tim", "Atış Puanı": 88, "Spor Puanı": 92, "Görev Adedi": 14},
        {"Personel": "Mehmet Kaya", "Tabur": "1. Tabur", "Bölük": "1. Bölük", "Tim": "1. Tim", "Atış Puanı": 95, "Spor Puanı": 85, "Görev Adedi": 18},
        {"Personel": "Ali Demir", "Tabur": "1. Tabur", "Bölük": "1. Bölük", "Tim": "2. Tim", "Atış Puanı": 76, "Spor Puanı": 80, "Görev Adedi": 10},
        {"Personel": "Mustafa Çelik", "Tabur": "1. Tabur", "Bölük": "2. Bölük", "Tim": "1. Tim", "Atış Puanı": 90, "Spor Puanı": 95, "Görev Adedi": 22},
        {"Personel": "Hüseyin Şahin", "Tabur": "1. Tabur", "Bölük": "2. Bölük", "Tim": "2. Tim", "Atış Puanı": 82, "Spor Puanı": 89, "Görev Adedi": 15},
    ])

# --- OTO-SUNUM (.PPTX) ÜRETİCİ FONKSİYON ---
def generate_pptx_presentation(filtered_df, filter_label):
    prs = Presentation()
    
    # Slayt 1: Başlık Slaydı
    blank_slide_layout = prs.slide_layouts[6]
    slide1 = prs.slides.add_slide(blank_slide_layout)
    
    # Başlık Kutusu
    txBox = slide1.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(2))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "PERFORMANS VE DURUM SUNUMU"
    p.font.bold = True
    p.font.size = Pt(36)
    p.font.color.rgb = RGBColor(24, 43, 73)
    p.alignment = PP_ALIGN.CENTER
    
    p2 = tf.add_paragraph()
    p2.text = f"Kapsam: {filter_label}\nTarih: {datetime.date.today().strftime('%d.%m.%Y')}"
    p2.font.size = Pt(18)
    p2.font.color.rgb = RGBColor(100, 100, 100)
    p2.alignment = PP_ALIGN.CENTER

    # Slayt 2: Grafik ve Görsel Analiz Slaydı
    slide2 = prs.slides.add_slide(blank_slide_layout)
    
    # Başlık
    title_box = slide2.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(9), Inches(0.8))
    tf2 = title_box.text_frame
    p3 = tf2.paragraphs[0]
    p3.text = f"Performans Analiz Grafiği ({filter_label})"
    p3.font.bold = True
    p3.font.size = Pt(22)
    p3.font.color.rgb = RGBColor(24, 43, 73)

    if not filtered_df.empty:
        # Matplotlib ile Profesyonel Çubuk / Bar Grafiği Oluşturma
        fig, ax = plt.subplots(figsize=(8, 4.5), dpi=200)
        
        x = np.arange(len(filtered_df["Personel"]))
        width = 0.35
        
        rects1 = ax.bar(x - width/2, filtered_df["Atış Puanı"], width, label='Atış Puanı', color='#1f77b4')
        rects2 = ax.bar(x + width/2, filtered_df["Spor Puanı"], width, label='Spor Puanı', color='#ff7f0e')
        
        ax.set_ylabel('Puanlar (0 - 100)')
        ax.set_title('Personel Bazlı Performans Karşılaştırması')
        ax.set_xticks(x)
        ax.set_xticklabels(filtered_df["Personel"], rotation=15, ha="right")
        ax.legend()
        ax.set_ylim(0, 110)
        ax.grid(axis='y', linestyle='--', alpha=0.7)
        
        plt.tight_layout()
        
        # Grafiği hafızadaki bir tampona kaydetme
        img_buf = io.BytesIO()
        plt.savefig(img_buf, format='png', bbox_inches='tight')
        img_buf.seek(0)
        plt.close(fig)
        
        # Görseli slayda yerleştirme
        slide2.shapes.add_picture(img_buf, Inches(0.8), Inches(1.5), width=Inches(8.4))

    # Slayt 3: Özet Tablo Slaydı
    slide3 = prs.slides.add_slide(blank_slide_layout)
    
    title_box3 = slide3.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(9), Inches(0.8))
    p4 = title_box3.text_frame.paragraphs[0]
    p4.text = "Detaylı Veri Listesi"
    p4.font.bold = True
    p4.font.size = Pt(22)
    p4.font.color.rgb = RGBColor(24, 43, 73)

    if not filtered_df.empty:
        rows, cols = filtered_df.shape[0] + 1, filtered_df.shape[1]
        left, top, width, height = Inches(0.5), Inches(1.5), Inches(9), Inches(0.5 * rows)
        table_shape = slide3.shapes.add_table(rows, cols, left, top, width, height)
        table = table_shape.table

        # Sütun Başlıkları
        for col_idx, col_name in enumerate(filtered_df.columns):
            cell = table.cell(0, col_idx)
            cell.text = str(col_name)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(24, 43, 73)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.color.rgb = RGBColor(255, 255, 255)
                paragraph.font.bold = True
                paragraph.font.size = Pt(11)

        # Tablo Verileri
        for row_idx, row_data in enumerate(filtered_df.values):
            for col_idx, value in enumerate(row_data):
                cell = table.cell(row_idx + 1, col_idx)
                cell.text = str(value)
                for paragraph in cell.text_frame.paragraphs:
                    paragraph.font.size = Pt(10)

    output = io.BytesIO()
    prs.save(output)
    output.seek(0)
    return output

# --- GİRİŞ EKRANI ---
if st.session_state["logged_in_user"] is None:
    st.title("🔐 Sistem Girişi")
    col1, col2 = st.columns([1, 2])
    with col1:
        username = st.text_input("Kullanıcı Adı")
        password = st.text_input("Şifre", type="password")
        if st.button("Giriş Yap"):
            if username in st.session_state["users_db"] and st.session_state["users_db"][username]["password"] == password:
                st.session_state["logged_in_user"] = username
                st.rerun()
            else:
                st.error("Hatalı kullanıcı adı veya şifre!")
    st.stop()

# --- OTURUM AÇILDIKTAN SONRAKİ ARAYÜZ ---
current_user = st.session_state["logged_in_user"]
user_role = st.session_state["users_db"][current_user]["role"]

st.sidebar.title(f"👤 {st.session_state['users_db'][current_user]['name']}")
st.sidebar.info(f"Yetki Rolü: **{user_role.upper()}**")

if st.sidebar.button("Çıkış Yap"):
    st.session_state["logged_in_user"] = None
    st.rerun()

# --- ROL VE İZİN KONTROLÜ İLE SEKMELERİN OLUŞTURULMASI ---
tabs = []
if user_role == "admin":
    # Admin Paneli: "Yeni Yönetici Ekle" kaldırıldı, sadece "Şifre Sıfırlama" bırakıldı.
    tabs = ["🔑 Şifre Sıfırlama", "📝 Veri Girişi", "📊 Sunum ve Rapor İndir"]
elif user_role == "reporter":
    # Reporter Yetkileri: Sadece tüm birimlere Veri Girişi ve Sunum/Rapor İndirme hakkı
    tabs = ["📝 Veri Girişi", "📊 Sunum ve Rapor İndir"]

active_tab = st.tabs(tabs)

tab_idx = 0

# --- SEKMELERİN İÇERİĞİ ---

# 1. ŞİFRE SIFIRLAMA SEKMESİ (Sadece Admin Görür)
if user_role == "admin":
    with active_tab[tab_idx]:
        st.subheader("🔑 Kullanıcı Şifre Sıfırlama Paneli")
        target_user = st.selectbox("Şifresi Sıfırlanacak Kullanıcı", list(st.session_state["users_db"].keys()))
        new_pass = st.text_input("Yeni Şifre", type="password")
        if st.button("Şifreyi Güncelle"):
            if new_pass:
                st.session_state["users_db"][target_user]["password"] = new_pass
                st.success(f"'{target_user}' kullanıcısının şifresi başarıyla güncellendi.")
            else:
                st.warning("Lütfen geçerli bir şifre girin.")
    tab_idx += 1

# 2. VERİ GİRİŞİ SEKMESİ (Reporter ve Admin Tüm Birimlere Veri Girebilir)
with active_tab[tab_idx]:
    st.subheader("📝 Tüm Birimler İçin Veri Giriş Ekranı")
    st.caption("Not: Reporter ve Admin rolündeki personel tüm tabur, bölük ve timler için veri girişi yapabilir.")
    
    with st.form("data_entry_form"):
        col1, col2, col3 = st.columns(3)
        with col1:
            personel_adi = st.text_input("Personel Adı Soyadı")
            tabur_adi = st.selectbox("Tabur", ["1. Tabur", "2. Tabur"])
        with col2:
            boluk_adi = st.selectbox("Bölük", ["1. Bölük", "2. Bölük", "3. Bölük"])
            tim_adi = st.selectbox("Tim", ["1. Tim", "2. Tim", "3. Tim"])
        with col3:
            atis_puani = st.number_input("Atış Puanı", min_value=0, max_value=100, value=85)
            spor_puani = st.number_input("Spor Puanı", min_value=0, max_value=100, value=85)
            gorev_adedi = st.number_input("Görev Adedi", min_value=0, value=10)
            
        submit_btn = st.form_submit_button("Veriyi Kaydet")
        if submit_btn:
            if personel_adi:
                new_row = {
                    "Personel": personel_adi,
                    "Tabur": tabur_adi,
                    "Bölük": boluk_adi,
                    "Tim": tim_adi,
                    "Atış Puanı": atis_puani,
                    "Spor Puanı": spor_puani,
                    "Görev Adedi": gorev_adedi
                }
                st.session_state["main_data"] = pd.concat([st.session_state["main_data"], pd.DataFrame([new_row])], ignore_index=True)
                st.success(f"{personel_adi} isimli personel verisi sisteme eklendi.")
            else:
                st.error("Lütfen personel adını giriniz.")

    st.markdown("---")
    st.write("### Mevcut Kayıtlar")
    st.dataframe(st.session_state["main_data"], use_container_width=True)

tab_idx += 1

# 3. SUNUM VE RAPOR İNDİR SEKMESİ (.PPTX İle Hazır Grafik Sunumu)
with active_tab[tab_idx]:
    st.subheader("📊 Profesyonel PowerPoint (.pptx) Sunum Oluşturucu")
    st.write("İndirmek istediğiniz veri kapsamını seçiniz. İndirilen sunum dosyası görseller ve grafiklerle birlikte sunuma hazır şekilde oluşacaktır.")

    # Filtreleme Seçenekleri
    col1, col2, col3, col4 = st.columns(4)
    
    df = st.session_state["main_data"]
    
    with col1:
        filter_type = st.radio("Kapsam Türü", ["Tüm Tabur", "Bölük Bazlı", "Tim Bazlı", "Personel Bazlı"])

    filtered_df = df.copy()
    filter_label = "Tüm Tabur"

    with col2:
        if filter_type == "Bölük Bazlı":
            selected_boluk = st.selectbox("Bölük Seçiniz", df["Bölük"].unique())
            filtered_df = df[df["Bölük"] == selected_boluk]
            filter_label = f"{selected_boluk}"
        elif filter_type == "Tim Bazlı":
            selected_boluk = st.selectbox("Bölük Seçiniz", df["Bölük"].unique())
            available_tims = df[df["Bölük"] == selected_boluk]["Tim"].unique()
            selected_tim = st.selectbox("Tim Seçiniz", available_tims)
            filtered_df = df[(df["Bölük"] == selected_boluk) & (df["Tim"] == selected_tim)]
            filter_label = f"{selected_boluk} - {selected_tim}"
        elif filter_type == "Personel Bazlı":
            selected_personel = st.selectbox("Personel Seçiniz", df["Personel"].unique())
            filtered_df = df[df["Personel"] == selected_personel]
            filter_label = f"Personel: {selected_personel}"

    st.markdown("---")
    st.write(f"**Seçilen Kapsam:** {filter_label} ({len(filtered_df)} Kayıt Bulundu)")
    st.dataframe(filtered_df, use_container_width=True)

    # Sunum Oluşturma ve İndirme Butonu
    if not filtered_df.empty:
        pptx_data = generate_pptx_presentation(filtered_df, filter_label)
        
        st.download_button(
            label="📈 PowerPoint Sunumunu Hazır (.pptx) Olarak İndir",
            data=pptx_data,
            file_name=f"Sunum_{filter_label.replace(' ', '_')}.pptx",
            mime="application/vnd.openxmlformats-officedocument.presentationml.presentation"
        )
    else:
        st.warning("Seçilen kriterlere uygun veri bulunamadı.")
